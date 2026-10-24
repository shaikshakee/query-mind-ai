import io, re
from pathlib import Path
from docx import Document
from pypdf import PdfReader

ALLOWED_EXTENSIONS = {".pdf",".docx",".txt",".md",".csv",".json",".py",".xml",".html",".htm",".log",".yaml",".yml"}

def extract_text(filename: str, raw: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError("Unsupported file format.")
    if ext == ".pdf":
        text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(raw)).pages)
    elif ext == ".docx":
        doc = Document(io.BytesIO(raw))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    else:
        text = raw.decode("utf-8", errors="replace")
    return text[:12000]

def document_context(question: str, state: dict) -> str:
    docs = state.get("documents", {})
    if not docs:
        return ""
    words = set(re.findall(r"\b[a-zA-Z0-9_]{3,}\b", question.lower()))
    ranked = []
    for filename, content in docs.items():
        lower = content.lower()
        score = sum(1 for word in words if word in lower)
        if score:
            ranked.append((score, filename, content))
    ranked.sort(key=lambda item: item[0], reverse=True)
    if not ranked and any(term in question.lower() for term in ("summarize","summary","uploaded document","uploaded file","the document","the file")):
        ranked = [(0, name, content) for name, content in list(docs.items())[:2]]
    return "\n\n".join(f"DOCUMENT: {name}\n{content[:3500]}" for _, name, content in ranked[:2])[:7000]
